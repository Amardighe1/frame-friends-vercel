from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont
import os

def create_placeholder_image(filename, text, size=(600, 400), color=(73, 109, 137)):
    img = Image.new('RGB', size, color=color)
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 30)
    except:
        font = ImageFont.load_default()
    
    text_bbox = d.textbbox((0, 0), text, font=font)
    text_width = text_bbox[2] - text_bbox[0]
    text_height = text_bbox[3] - text_bbox[1]
    
    x = (size[0] - text_width) / 2
    y = (size[1] - text_height) / 2
    
    d.text((x, y), text, fill=(255, 255, 255), font=font)
    img.save(filename)

def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level)
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = None

def main():
    doc = Document()
    
    # Title Page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("A Micro-Project Report ON\n\n")
    run.font.size = Pt(14)
    run.font.bold = True
    
    title = p.add_run("FarmFriend: AI Crop Disease Detector, Emergency Vet & Market Prices\n\n")
    title.font.size = Pt(22)
    title.font.bold = True
    
    sub = p.add_run("Submitted in partial fulfillment of the requirements for the Diploma in Computer Engineering\n")
    sub.font.size = Pt(12)
    
    run2 = p.add_run("\nSubject: Mobile Application Development (316006)")
    run2.font.size = Pt(14)
    run2.font.bold = True
    
    doc.add_page_break()
    
    # Generating Images
    create_placeholder_image("splash_screen.png", "FarmFriend App Splash Screen")
    create_placeholder_image("code_snippet.png", "Android Studio Code Screenshot\n(Camera Intent & UI Layout)", size=(700, 400), color=(43, 43, 43))
    create_placeholder_image("virtual_device.png", "Virtual Running Image (FarmFriend Home)", size=(300, 600), color=(100, 200, 100))
    create_placeholder_image("camera_detection.png", "AI Crop Disease Detection\nRunning on Emulator", size=(300, 600), color=(50, 150, 200))

    # Index
    add_heading(doc, "INDEX", level=1)
    index_content = [
        "1. Rationale",
        "2. Aims/Benefits of the Micro-Project",
        "3. Course Outcomes Addressed",
        "4. Proposed Methodology",
        "5. System Architecture & Features",
        "6. Android Application Code (UI & Backend)",
        "7. Project Execution UI and Running App Images",
        "8. Tools and Technologies Used",
        "9. Skill Developed",
        "10. Conclusion & Future Scope"
    ]
    for idx, item in enumerate(index_content):
        doc.add_paragraph(item)
    
    doc.add_page_break()

    # 1. Rationale
    add_heading(doc, "1. Rationale")
    doc.add_paragraph(
        "Android OS is one of the fastest-growing environments worldwide. In the context of India's agricultural ecosystem, "
        "mobile applications have the potential to solve critical challenges faced by farmers natively. The proposed application, "
        "FarmFriend, aims to bridge the gap between technology and agriculture. Current statistics indicate that farmers often lack "
        "accurate guidance on crop diseases, direct access to immediate veterinary care, and transparent market pricing. "
        "This project builds an Android app using components from Android SDK to deliver real-time assistance through AI-based diagnostic "
        "scanning and database retrieval APIs."
    )
    
    # 2. Aims/Benefits
    add_heading(doc, "2. Aims/Benefits of the Micro-Project")
    doc.add_paragraph(
        "• To design graphical user interfaces using standard Android UI Layouts (Constraint Layout, Linear Layout).\n"
        "• To capture and process images using Camera Intents for AI disease detection.\n"
        "• To utilize API communication to fetch dynamic current market prices.\n"
        "• To construct local database queries using SQLite/ Firebase for emergency vet contacts.\n"
        "• To package and test an Android environment app targeting practical real-world usage."
    )

    # 3. Course Outcomes
    add_heading(doc, "3. Course Outcomes Addressed")
    doc.add_paragraph(
        "The project maps strongly with the following COs of Mobile Application Development (316006):\n"
        "CO1 - Interpret the features of Android Operating System.\n"
        "CO2 - Use Android configuration and development environment tools.\n"
        "CO3 - Develop Android Applications using UI Components and Layouts.\n"
        "CO4 - Create database-driven Android Applications (Veterinary Contacts).\n"
        "CO5 - Develop Android applications that require relevant permissions (Camera, Internet, Storage)."
    )

    # 4. Proposed Methodology
    add_heading(doc, "4. Proposed Methodology")
    doc.add_paragraph(
        "1. Problem Definition: Analyzed the requirements for the FarmFriend app.\n"
        "2. UI/UX Design: Prototyped the application using Android Studio’s Layout Editor (XML). Integrated RecyclerViews for Market Prices.\n"
        "3. Core Components implementation: Created explicit and implicit intents to trigger camera scanning.\n"
        "4. Backend Connectivity: Formatted REST API endpoints dynamically reading JSON to list local market prices dynamically.\n"
        "5. Testing: Built the debug APK and ran the application on AVD (Android Virtual Device) emphasizing logging via Logcat."
    )

    # 5. System Architecture & Features
    add_heading(doc, "5. System Architecture & Features")
    doc.add_paragraph(
        "FarmFriend covers three distinct agricultural pillars in its navigation flow:\n"
        "- AI Crop Disease Detector: Users click a photo of their damaged crop leaf. The application converts the bitmap, passes it "
        "through an inference model (e.g. TF Lite) or posts it to a server, and suggests pesticide steps.\n"
        "- Emergency Vet Locator: Using location services, it queries a Firebase real-time database to retrieve closest veterinarians.\n"
        "- Mandi (Market) Prices: Consumes an external APMC API fetching live crop pricing based on the selected region."
    )
    doc.add_picture("splash_screen.png", width=Inches(4))

    # 6. Android Code
    add_heading(doc, "6. Android Application Code (UI & Backend)")
    doc.add_paragraph("XML Layout for Homepage (activity_main.xml):")
    doc.add_paragraph(
        '''
<LinearLayout xmlns:android="http://schemas.android.com/apk/res/android"
    android:layout_width="match_parent"
    android:layout_height="match_parent"
    android:orientation="vertical"
    android:padding="16dp">

    <TextView
        android:id="@+id/tvWelcome"
        android:text="Welcome to FarmFriend"
        android:textSize="24sp"
        android:textStyle="bold"
        android:layout_width="wrap_content"
        android:layout_height="wrap_content" />

    <Button
        android:id="@+id/btnScanCrop"
        android:text="AI Crop Disease Detector"
        android:layout_marginTop="20dp"
        android:layout_width="match_parent"
        android:layout_height="wrap_content" />
</LinearLayout>
        '''
    , style='Intense Quote')

    doc.add_paragraph("Java/Kotlin Integration for Camera Intent (MainActivity.java):")
    doc.add_paragraph(
        '''
Intent takePictureIntent = new Intent(MediaStore.ACTION_IMAGE_CAPTURE);
if (takePictureIntent.resolveActivity(getPackageManager()) != null) {
    startActivityForResult(takePictureIntent, REQUEST_IMAGE_CAPTURE);
}
        '''
    , style='Intense Quote')
    doc.add_paragraph("Below is the Android Studio code execution interface (IDE Overview).")
    doc.add_picture("code_snippet.png", width=Inches(6))

    doc.add_page_break()

    # 7. Project Execution UI and Running App Images
    add_heading(doc, "7. Project Execution UI and Running App Images")
    doc.add_paragraph("This section showcases the running state of the FarmFriend application inside the Android Virtual Device (AVD).")
    doc.add_picture("virtual_device.png", width=Inches(2.5))
    doc.add_picture("camera_detection.png", width=Inches(2.5))

    # 8. Tools and Technologies Used
    add_heading(doc, "8. Tools and Technologies Used")
    doc.add_paragraph(
        "- IDE: Android Studio Giraffe & Above\n"
        "- Language: Java / XML\n"
        "- Deployment: Dalvik Virtual Machine (DVM) / Android Runtime (ART)\n"
        "- Services: Firebase Database, Android Implicit Intents, Location APIs\n"
        "- Hardware: PC with 8GB+ RAM\n"
        "- Version Control: Git / GitHub"
    )

    # 9. Skill Developed
    add_heading(doc, "9. Skill Developed")
    doc.add_paragraph(
        "- Developed the ability to structure complex XML layout hierarchies effectively.\n"
        "- Mastered API consumptions and JSON Parsing within Android.\n"
        "- Learned App-level permission management (Camera, Internet, GPS).\n"
        "- Improved hardware-accelerated emulation configuration and application debugging."
    )

    # 10. Conclusion & Future Scope
    add_heading(doc, "10. Conclusion & Future Scope")
    doc.add_paragraph(
        "Conclusion:\n"
        "The FarmFriend Android application successfully implements real-world features required for the Indian agricultural base. "
        "Through applying theoretical Mobile Application Development (316006) concepts into practice, we demonstrated effective UI component usage, hardware interactions, and network programming natively on Android.\n\n"
        "Future Scope:\n"
        "- Implementing local ML on-device mapping without Internet dependency.\n"
        "- Real-time push notifications indicating abrupt weather changes or Mandi price jumps via Google Cloud Messaging / Firebase."
    )

    doc.save("FarmFriend_MicroProject_Report.docx")
    
    # Cleanup temporary images
    for file in ["splash_screen.png", "code_snippet.png", "virtual_device.png", "camera_detection.png"]:
        if os.path.exists(file):
            os.remove(file)

if __name__ == "__main__":
    main()
